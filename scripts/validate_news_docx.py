#!/usr/bin/env python
"""Validate a News_SciPaper DOCX output before delivery."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from docx import Document


REQUIRED_FIELDS = ["项目支持：", "论文链接：", "在线发表日期：", "影响因子："]
DEFAULT_MIN_CHARS = 1000
DEFAULT_MAX_CHARS = 1300
DEFAULT_MIN_IMAGES = 2
DEFAULT_MAX_IMAGES = 4
DEFAULT_IMAGE_WIDTH_CM = 14.65


def read_docx(path: Path) -> tuple[Document, list[str]]:
    doc = Document(str(path))
    texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    return doc, texts


def main_body_texts(texts: list[str]) -> list[str]:
    if not texts:
        return []
    body: list[str] = []
    for text in texts[1:]:
        if any(text.startswith(field) for field in REQUIRED_FIELDS):
            break
        if re.match(r"^图\d+", text):
            break
        if text.startswith(("论文链接：", "文章链接：")):
            break
        body.append(text)
    return body


def validate(
    path: Path,
    min_chars: int,
    max_chars: int,
    min_images: int,
    max_images: int,
    expected_width_cm: float,
    width_tolerance_cm: float,
) -> dict[str, Any]:
    doc, texts = read_docx(path)
    body = main_body_texts(texts)
    body_chars = sum(len(text) for text in body)
    image_widths = [round(shape.width.cm, 2) for shape in doc.inline_shapes]
    full_text = "\n".join(texts)

    checks = {
        "body_length_ok": min_chars <= body_chars <= max_chars,
        "body_paragraph_count_ok": 4 <= len(body) <= 6,
        "required_fields_ok": {field: any(text.startswith(field) for text in texts) for field in REQUIRED_FIELDS},
        "image_count_ok": min_images <= len(doc.inline_shapes) <= max_images,
        "image_width_ok": all(abs(width - expected_width_cm) <= width_tolerance_cm for width in image_widths),
        "no_question_mark_garbled_text": "?" not in full_text,
        "no_unconfirmed_placeholders": not any(marker in full_text for marker in ["【请确认", "【待", "待插图", "待确认中文"]),
        "figure_captions_ok": len([text for text in texts if re.match(r"^图\d+", text)]) >= len(doc.inline_shapes),
    }

    errors = []
    warnings = []
    if not checks["body_length_ok"]:
        errors.append(f"正文主体长度为 {body_chars} 字，不在 {min_chars}-{max_chars} 字范围。")
    if not checks["body_paragraph_count_ok"]:
        warnings.append(f"正文主体为 {len(body)} 段，建议为 4-6 段。")
    missing_fields = [field for field, ok in checks["required_fields_ok"].items() if not ok]
    if missing_fields:
        errors.append("缺少固定字段：" + "、".join(missing_fields))
    if not checks["image_count_ok"]:
        errors.append(f"图片数量为 {len(doc.inline_shapes)} 张，不在 {min_images}-{max_images} 张范围。")
    if not checks["image_width_ok"]:
        errors.append(f"图片宽度为 {image_widths} cm，未全部接近 {expected_width_cm} cm。")
    if not checks["no_question_mark_garbled_text"]:
        errors.append("正文中存在问号，可能有编码损坏或未替换占位。")
    if not checks["no_unconfirmed_placeholders"]:
        errors.append("正文中存在未确认占位符。")
    if not checks["figure_captions_ok"]:
        errors.append("图片数量多于图题数量，可能缺少中文图题。")

    return {
        "path": str(path),
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "metrics": {
            "title": texts[0] if texts else "",
            "body_chars": body_chars,
            "body_paragraphs": len(body),
            "paragraphs_total": len(texts),
            "image_count": len(doc.inline_shapes),
            "image_widths_cm": image_widths,
            "figure_caption_count": len([text for text in texts if re.match(r"^图\d+", text)]),
        },
        "checks": checks,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--min-chars", type=int, default=DEFAULT_MIN_CHARS)
    parser.add_argument("--max-chars", type=int, default=DEFAULT_MAX_CHARS)
    parser.add_argument("--min-images", type=int, default=DEFAULT_MIN_IMAGES)
    parser.add_argument("--max-images", type=int, default=DEFAULT_MAX_IMAGES)
    parser.add_argument("--expected-image-width-cm", type=float, default=DEFAULT_IMAGE_WIDTH_CM)
    parser.add_argument("--width-tolerance-cm", type=float, default=0.05)
    parser.add_argument("--out", type=Path, help="Optional JSON output path.")
    args = parser.parse_args()

    result = validate(
        args.docx,
        args.min_chars,
        args.max_chars,
        args.min_images,
        args.max_images,
        args.expected_image_width_cm,
        args.width_tolerance_cm,
    )
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(text + "\n", encoding="utf-8")
    print(text)
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
